#!/usr/bin/env python3
"""Dump what is in a .pptx as Markdown: slide size, layouts, and per slide the
layout name, every shape (type, placeholder role, position in inches) and its
text, tables, chart types, pictures and speaker notes.

Usage:
  pptx_outline.py deck.pptx                 # full outline
  pptx_outline.py deck.pptx --text-only     # just titles/text/notes (for reading content)
  pptx_outline.py deck.pptx --layouts       # list masters/layouts with their placeholders (template analysis)
  pptx_outline.py deck.pptx --slides 2,5-7
"""
import argparse, sys
from pptx import Presentation
from pptx.util import Emu

IN = 914400


def rng(spec, n):
    if not spec:
        return list(range(1, n + 1))
    out = []
    for part in spec.split(","):
        if "-" in part:
            a, b = part.split("-")
            out += range(int(a), int(b) + 1)
        else:
            out.append(int(part))
    return [i for i in out if 1 <= i <= n]


def pos(sh):
    try:
        return f"@({sh.left / IN:.2f},{sh.top / IN:.2f}) {sh.width / IN:.2f}x{sh.height / IN:.2f}in"
    except (TypeError, AttributeError):
        return "@(inherited)"


def kind(sh):
    if sh.is_placeholder:
        return f"placeholder:{str(sh.placeholder_format.type).split('.')[-1].split(' ')[0].lower()}#{sh.placeholder_format.idx}"
    if sh.has_chart if hasattr(sh, "has_chart") else False:
        return f"chart:{str(sh.chart.chart_type).split('.')[-1].split(' ')[0].lower()}"
    if getattr(sh, "has_table", False) and sh.has_table:
        return "table"
    return str(sh.shape_type).split(".")[-1].split(" ")[0].lower() if sh.shape_type else "shape"


def shape_lines(sh, text_only, depth=0):
    pad = "  " * depth
    out = []
    if sh.shape_type is not None and str(sh.shape_type).startswith("GROUP"):
        if not text_only:
            out.append(f"{pad}- group '{sh.name}' {pos(sh)}")
        for c in sh.shapes:
            out += shape_lines(c, text_only, depth + 1)
        return out
    head = f"{pad}- {kind(sh)} '{sh.name}' {pos(sh)}"
    body = []
    if sh.has_text_frame and sh.text_frame.text.strip():
        for p in sh.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t:
                body.append(f"{pad}    {'  ' * (p.level or 0)}{t}")
    if getattr(sh, "has_table", False) and sh.has_table:
        for row in sh.table.rows:
            body.append(f"{pad}    | " + " | ".join(c.text.replace("\n", " ") for c in row.cells) + " |")
    if getattr(sh, "has_chart", False) and sh.has_chart:
        ch = sh.chart
        try:
            cats = list(ch.plots[0].categories)
            body.append(f"{pad}    categories: {cats}")
            for s in ch.plots[0].series:
                body.append(f"{pad}    series {s.name!r}: {list(s.values)}")
        except Exception:
            pass
    if kind(sh) == "picture":
        try:
            body.append(f"{pad}    image: {sh.image.filename} {sh.image.size[0]}x{sh.image.size[1]}px")
        except Exception:
            pass
    if text_only:
        return [b.strip() if not b.strip().startswith("|") else b.strip() for b in body]
    return [head] + body


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("--text-only", action="store_true")
    ap.add_argument("--layouts", action="store_true")
    ap.add_argument("--slides")
    a = ap.parse_args()
    prs = Presentation(a.deck)
    print(f"# {a.deck}\n\nslide size {prs.slide_width / IN:.3f} x {prs.slide_height / IN:.3f} in, {len(prs.slides)} slides\n")
    if a.layouts:
        for mi, m in enumerate(prs.slide_masters):
            print(f"## master {mi + 1}")
            for li, lay in enumerate(m.slide_layouts):
                used = sum(1 for s in prs.slides if s.slide_layout == lay)
                phs = ", ".join(f"{str(p.placeholder_format.type).split('.')[-1].split(' ')[0].lower()}#{p.placeholder_format.idx}" for p in lay.placeholders)
                print(f"- [{li}] {lay.name!r} used by {used} slide(s): {phs or 'no placeholders'}")
        return
    for i in rng(a.slides, len(prs.slides)):
        s = prs.slides[i - 1]
        hidden = s._element.get("show") == "0"
        print(f"## Slide {i}{' (hidden)' if hidden else ''} — layout {s.slide_layout.name!r}")
        for sh in s.shapes:
            for line in shape_lines(sh, a.text_only):
                print(line)
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            print("  notes: " + s.notes_slide.notes_text_frame.text.strip().replace("\n", " / "))
        print()


if __name__ == "__main__":
    main()
