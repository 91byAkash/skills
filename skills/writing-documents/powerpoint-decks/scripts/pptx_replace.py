#!/usr/bin/env python3
"""Replace text in a deck while keeping each paragraph's formatting — the safe
way to fill a template's slots. Matches across run boundaries (PowerPoint often
splits one word over several runs).

Usage:
  pptx_replace.py deck.pptx -o out.pptx --find "{{client}}" --replace "Acme Ltd"
  pptx_replace.py deck.pptx -o out.pptx --pairs pairs.json        # {"old": "new", ...}
  pptx_replace.py deck.pptx -o out.pptx --pairs pairs.json --slides 2-4 --notes
  pptx_replace.py deck.pptx --list-tokens                          # show {{...}} / [..] tokens left

The replacement text takes the formatting of the run where the match starts;
runs fully covered by the match are emptied, partly covered runs keep their
remainder, so bold labels such as "Owner:" survive when the match is after them.
Searches text boxes, placeholders, grouped shapes, table cells and (with
--notes) speaker notes. Exit 2 if any --find/pairs key was not found.
"""
import argparse, json, re, sys
from pptx import Presentation

TOKEN_RE = re.compile(r"\{\{[^}]+\}\}|\[[A-Z][^\]]{1,40}\]|<<[^>]+>>")


def frames(shapes):
    for sh in shapes:
        if sh.shape_type is not None and str(sh.shape_type).startswith("GROUP"):
            yield from frames(sh.shapes)
            continue
        if sh.has_text_frame:
            yield sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    yield c.text_frame


def replace_in_paragraph(p, old, new):
    runs = list(p.runs)
    if not runs:
        return 0
    count = 0
    while True:
        full = "".join(r.text for r in runs)
        i = full.find(old)
        if i < 0:
            return count
        j = i + len(old)
        pos = 0
        first = True
        for r in runs:
            a, b = pos, pos + len(r.text)
            pos = b
            if b <= i or a >= j:
                continue
            lo, hi = max(i, a) - a, min(j, b) - a
            if first:
                r.text = r.text[:lo] + new + r.text[hi:]
                first = False
            else:
                r.text = r.text[:lo] + r.text[hi:]
        count += 1
        if old in new:  # avoid infinite loop when the replacement contains the search text
            return count


def rng(spec, n):
    if not spec:
        return set(range(1, n + 1))
    out = set()
    for part in spec.split(","):
        lo, _, hi = part.partition("-")
        out.update(range(int(lo), int(hi or lo) + 1))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("-o", "--out")
    ap.add_argument("--find")
    ap.add_argument("--replace", default="")
    ap.add_argument("--pairs")
    ap.add_argument("--slides")
    ap.add_argument("--notes", action="store_true")
    ap.add_argument("--list-tokens", action="store_true")
    a = ap.parse_args()
    prs = Presentation(a.deck)
    want = rng(a.slides, len(prs.slides))
    if a.list_tokens:
        found = 0
        for i, s in enumerate(prs.slides, 1):
            for tf in frames(s.shapes):
                for t in TOKEN_RE.findall(tf.text):
                    print(f"slide {i}: {t}")
                    found += 1
        print(f"{found} token(s)")
        return
    pairs = {}
    if a.pairs:
        pairs.update(json.load(open(a.pairs, encoding="utf-8")))
    if a.find:
        pairs[a.find] = a.replace
    if not pairs or not a.out:
        sys.exit("need -o and --find/--replace or --pairs")
    hits = {k: 0 for k in pairs}
    for i, s in enumerate(prs.slides, 1):
        if i not in want:
            continue
        tfs = list(frames(s.shapes))
        if a.notes and s.has_notes_slide:
            tfs.append(s.notes_slide.notes_text_frame)
        for tf in tfs:
            for p in tf.paragraphs:
                for old, new in pairs.items():
                    hits[old] += replace_in_paragraph(p, old, new)
    prs.save(a.out)
    for k, v in hits.items():
        print(f"{v:>4} x {k!r}")
    missing = [k for k, v in hits.items() if v == 0]
    if missing:
        print(f"not found: {missing}", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()
