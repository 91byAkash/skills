#!/usr/bin/env python3
"""Generate a small, convention-following projection model to copy from:
Assumptions -> Model -> Summary, with live formulas, colour-coded inputs,
named ranges, a scenario dropdown, number formats, a chart and source notes.

Usage:
  model_template.py -o model.xlsx [--years 5] [--start-year 2026] [--currency "$"] [--unit "k"]

Then:  recalc.py model.xlsx && xlsx_audit.py model.xlsx
Edit the ASSUMPTIONS list below (label, value, format, note) to fit the task;
every formula refers to those cells, so changing an input moves the whole model.
"""
import argparse

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEADER = PatternFill("solid", fgColor="1F3864")
THIN = Side(style="thin", color="9AA5B8")
FONT = "Arial"

# label, name, value, number format, source note
ASSUMPTIONS = [
    ("Starting revenue", "rev0", 1200, "money", "Source: FY2025 actuals provided by the user"),
    ("Revenue growth - Base", "g_base", 0.18, "pct", "Source: user estimate"),
    ("Revenue growth - Downside", "g_down", 0.08, "pct", "Source: user estimate"),
    ("Revenue growth - Upside", "g_up", 0.28, "pct", "Source: user estimate"),
    ("Gross margin", "gm", 0.62, "pct", "Source: trailing 12-month average"),
    ("Operating expenses (% of revenue)", "opex_pct", 0.41, "pct", "Assumption: held flat"),
    ("Tax rate", "tax", 0.25, "pct", "Assumption: statutory rate"),
]


def fmt(kind, cur, unit):
    return {"money": f'{cur}#,##0;({cur}#,##0);"-"', "pct": '0.0%;(0.0%);"-"', "mult": "0.0x", "int": "#,##0"}[kind]


def style_header(row):
    for c in row:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEADER
        c.alignment = Alignment(horizontal="center")


def build(out, years, start, cur, unit):
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "Assumptions"
    a["A1"] = "Assumptions"; a["A1"].font = Font(name=FONT, bold=True, size=14)
    a["A2"] = f"Blue = input you may change; yellow fill = key assumption. Currency {cur}{unit}."
    a["A2"].font = Font(name=FONT, italic=True, color="555555")
    a.append([])
    a.append(["Input", "Value", "Notes"]); style_header(a[4])
    row = 5
    for label, name, value, kind, note in ASSUMPTIONS:
        a.cell(row, 1, label).font = Font(name=FONT)
        c = a.cell(row, 2, value)
        c.font = Font(name=FONT, color=BLUE)
        c.number_format = fmt(kind, cur, unit)
        if name.startswith("g_"):
            c.fill = YELLOW
        c.comment = Comment(note, "model")
        a.cell(row, 3, note).font = Font(name=FONT, color="555555")
        wb.defined_names[name] = DefinedName(name, attr_text=f"Assumptions!$B${row}")
        row += 1
    row += 1
    a.cell(row, 1, "Scenario").font = Font(name=FONT, bold=True)
    sc = a.cell(row, 2, "Base"); sc.font = Font(name=FONT, color=BLUE); sc.fill = YELLOW
    dv = DataValidation(type="list", formula1='"Base,Downside,Upside"', allow_blank=False, showErrorMessage=True,
                        errorTitle="Scenario", error="Pick Base, Downside or Upside")
    a.add_data_validation(dv); dv.add(sc)
    wb.defined_names["scenario"] = DefinedName("scenario", attr_text=f"Assumptions!$B${row}")
    row += 1
    a.cell(row, 1, "Growth used").font = Font(name=FONT)
    g = a.cell(row, 2, '=IF(scenario="Upside",g_up,IF(scenario="Downside",g_down,g_base))')
    g.number_format = fmt("pct", cur, unit); g.font = Font(name=FONT, color=BLACK)
    wb.defined_names["growth"] = DefinedName("growth", attr_text=f"Assumptions!$B${row}")
    a.column_dimensions["A"].width = 36; a.column_dimensions["B"].width = 14; a.column_dimensions["C"].width = 48

    m = wb.create_sheet("Model")
    m["A1"] = f"Projection ({cur}{unit})"; m["A1"].font = Font(name=FONT, bold=True, size=14)
    m.append([])
    m.append(["Line item"] + [str(start + i) for i in range(years)])  # years as text, not 2,026
    style_header(m[3])
    lines = ["Revenue", "Gross profit", "Operating expenses", "Operating income", "Tax", "Net income", "Net margin"]
    for i, name in enumerate(lines):
        m.cell(4 + i, 1, name).font = Font(name=FONT, bold=name in ("Revenue", "Net income"))
    for j in range(years):
        col = get_column_letter(2 + j)
        prev = get_column_letter(1 + j)
        rev = "=rev0*(1+growth)" if j == 0 else f"={prev}4*(1+growth)"
        cells = {
            4: rev,
            5: f"={col}4*gm",
            6: f"=-{col}4*opex_pct",
            7: f"={col}5+{col}6",
            8: f"=-MAX({col}7,0)*tax",          # no tax credit on losses in this simple model
            9: f"={col}7+{col}8",
            10: f"=IF({col}4=0,0,{col}9/{col}4)",  # guard the denominator
        }
        for r, f in cells.items():
            c = m.cell(r, 2 + j, f)
            c.font = Font(name=FONT, color=GREEN if j == 0 and r == 4 else BLACK)
            c.number_format = fmt("pct" if r == 10 else "money", cur, unit)
            c.border = Border(top=THIN) if r in (7, 9) else Border()
        m.column_dimensions[col].width = 12
    m.column_dimensions["A"].width = 24
    m.freeze_panes = "B4"
    last = get_column_letter(1 + years)
    m.conditional_formatting.add(f"B9:{last}9", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000")))

    s = wb.create_sheet("Summary")
    s["A1"] = "Summary"; s["A1"].font = Font(name=FONT, bold=True, size=14)
    rows = [
        ("Scenario", "=scenario", None),
        (f"Revenue {start + years - 1}", f"=Model!{last}4", "money"),
        (f"Revenue CAGR {start}-{start + years - 1}", f"=IFERROR((Model!{last}4/Model!B4)^(1/(COLUMNS(Model!B4:{last}4)-1))-1,0)", "pct"),
        (f"Cumulative net income", f"=SUM(Model!B9:{last}9)", "money"),
        ("Average net margin", f"=AVERAGE(Model!B10:{last}10)", "pct"),
    ]
    for i, (label, f, kind) in enumerate(rows, start=3):
        s.cell(i, 1, label).font = Font(name=FONT)
        c = s.cell(i, 2, f); c.font = Font(name=FONT, color=GREEN)
        if kind:
            c.number_format = fmt(kind, cur, unit)
    s.column_dimensions["A"].width = 30; s.column_dimensions["B"].width = 16
    chart = BarChart(); chart.title = "Revenue and net income"; chart.y_axis.title = f"{cur}{unit}"
    chart.add_data(Reference(m, min_col=1, max_col=1 + years, min_row=4, max_row=4), from_rows=True, titles_from_data=True)
    chart.add_data(Reference(m, min_col=1, max_col=1 + years, min_row=9, max_row=9), from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(m, min_col=2, max_col=1 + years, min_row=3, max_row=3))
    chart.height, chart.width = 7.5, 16
    s.add_chart(chart, "D3")
    wb.calculation.fullCalcOnLoad = True  # Excel recomputes everything on open
    wb.save(out)
    print(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--years", type=int, default=5)
    ap.add_argument("--start-year", type=int, default=2026)
    ap.add_argument("--currency", default="$")
    ap.add_argument("--unit", default="k")
    a = ap.parse_args()
    build(a.out, a.years, a.start_year, a.currency, a.unit)


if __name__ == "__main__":
    main()
