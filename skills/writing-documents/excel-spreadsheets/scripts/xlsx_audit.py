#!/usr/bin/env python3
"""Static audit of a workbook's formulas and structure - no LibreOffice needed.
Run it before and after editing; pair it with recalc.py.

Usage:
  xlsx_audit.py model.xlsx [--json] [--sheet Model]

ERRORS (exit 1):
  - dynamic-array/new functions (XLOOKUP, FILTER, UNIQUE, SORT, SEQUENCE, LET, ...): openpyxl writes
    them without _xlfn./spill metadata and LibreOffice < 24.8 (< 25.8 for VSTACK, TEXTSPLIT...) can't evaluate them
  - post-2007 functions written without the _xlfn. prefix (TEXTJOIN, IFS, MAXIFS, ...)
  - references to sheets whose names need quotes but have none ('Q1 Inputs'!B2)
  - references to sheets that don't exist; defined names pointing at #REF!
  - cached error values (#DIV/0!, #REF!, ...) from the last calculation
WARNINGS:
  - a formula that breaks the pattern of its row/column neighbours
  - SUM/AVERAGE ranges that stop short of numbers directly above the total
  - hard-coded constants inside formulas (=B5*1.05)
  - hard-coded numbers typed over one cell of a formula column
  - percentage-formatted constants greater than 1 (15 shown as 1500%)
  - external workbook links, hidden sheets, merged cells, formulas with no cached value
"""
import argparse, json, re, sys, zipfile
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

UNSUPPORTED = {"XLOOKUP", "XMATCH", "FILTER", "UNIQUE", "SORT", "SORTBY", "SEQUENCE", "RANDARRAY", "LET", "LAMBDA",
               "TEXTSPLIT", "TEXTBEFORE", "TEXTAFTER", "VSTACK", "HSTACK", "TAKE", "DROP", "CHOOSECOLS", "CHOOSEROWS",
               "TOCOL", "TOROW", "WRAPROWS", "WRAPCOLS", "EXPAND", "MAP", "REDUCE", "SCAN", "BYROW", "BYCOL", "MAKEARRAY",
               "GROUPBY", "PIVOTBY"}
NEEDS_PREFIX = {"TEXTJOIN", "CONCAT", "IFS", "SWITCH", "MAXIFS", "MINIFS", "IFNA", "XOR", "DAYS", "ISOWEEKNUM",
                "STDEV.S", "STDEV.P", "VAR.S", "VAR.P", "PERCENTILE.INC", "PERCENTILE.EXC", "QUARTILE.INC",
                "RANK.EQ", "RANK.AVG", "CEILING.MATH", "FLOOR.MATH", "FORECAST.LINEAR", "NORM.DIST", "NORM.INV",
                "T.TEST", "CONFIDENCE.NORM", "MODE.SNGL", "AGGREGATE", "NUMBERVALUE", "UNICHAR", "UNICODE", "SHEET", "SHEETS"}
FUNC_RE = re.compile(r"(?<![A-Za-z0-9_.])((?:_xlfn\.|_xlws\.)?[A-Z][A-Z0-9.]*)\s*\(")
REF_RE = re.compile(r"(?<![A-Za-z0-9_!:$'])(\$?)([A-Z]{1,3})(\$?)(\d{1,7})(?![A-Za-z0-9_(])")
SHEETREF_RE = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_.]+))!")
CONST_RE = re.compile(r"[*/+\-^]\s*(\d+\.?\d*)(?![\d.:A-Za-z(])")
BENIGN_CONSTANTS = {"0", "1", "2", "4", "7", "10", "12", "24", "52", "60", "100", "365", "1000"}
ERRORS = ("#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#SPILL!", "#CALC!")


def strip_strings(f):
    return re.sub(r'"(?:[^"]|"")*"', '""', f)


def signature(formula, row, col):
    """Formula with relative refs expressed as offsets, so =B2*C2 in row 2 and
    =B3*C3 in row 3 compare equal."""
    def rel(m):
        cabs, cl, rabs, rw = m.groups()
        c = cl if cabs else f"C[{column_index_from_string(cl) - col}]"
        r = rw if rabs else f"R[{int(rw) - row}]"
        return f"{c}{r}"
    return REF_RE.sub(rel, strip_strings(formula))


def audit(path, only_sheet=None):
    errors, warnings, info = [], [], []
    keep_vba = str(path).lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(path, keep_vba=keep_vba)
    wbv = openpyxl.load_workbook(path, data_only=True)
    names = set(wb.sheetnames)
    with zipfile.ZipFile(path) as z:
        ext = [n for n in z.namelist() if n.startswith("xl/externalLinks/") and n.endswith(".xml")]
    if ext:
        warnings.append(f"{len(ext)} external workbook link(s): values depend on other files; re-saving + recalculating can break them")
    try:
        defined = list(wb.defined_names.values())          # openpyxl >= 3.1
    except AttributeError:
        defined = list(wb.defined_names.definedName)       # older openpyxl
    for ws_ in wb.worksheets:
        defined += list(getattr(ws_, "defined_names", {}).values()) if hasattr(getattr(ws_, "defined_names", None), "values") else []
    for dn in defined:
        if "#REF!" in str(dn.attr_text):
            errors.append(f"defined name {dn.name!r} points at #REF!")
    for ws in wb.worksheets:
        if only_sheet and ws.title != only_sheet:
            continue
        if ws.sheet_state != "visible":
            info.append(f"sheet {ws.title!r} is {ws.sheet_state}")
        if ws.merged_cells.ranges:
            info.append(f"{ws.title}: {len(ws.merged_cells.ranges)} merged range(s) - write only to the top-left cell")
        vs = wbv[ws.title]
        formulas, no_cache = {}, 0
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                loc = f"{ws.title}!{c.coordinate}"
                cached = vs[c.coordinate].value
                if isinstance(cached, str) and cached.startswith(ERRORS):
                    errors.append(f"{loc}: cached error {cached}")
                if isinstance(v, (int, float)) and not isinstance(v, bool) and "%" in (c.number_format or "") and abs(v) > 1:
                    warnings.append(f"{loc}: {v} formatted as percent shows {v * 100:g}% - store 0.15 for 15%")
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                formulas[(c.row, c.column)] = v
                if cached is None:
                    no_cache += 1
                body = strip_strings(v)
                for fn in FUNC_RE.findall(body):
                    bare = fn.replace("_xlfn.", "").replace("_xlws.", "")
                    if bare in UNSUPPORTED:
                        errors.append(f"{loc}: {bare} - needs _xlfn. + spill metadata openpyxl can't write, and LibreOffice >= 24.8/25.8 to verify; use INDEX/MATCH or compute in Python")
                    elif bare in NEEDS_PREFIX and not fn.startswith("_xlfn."):
                        errors.append(f"{loc}: write _xlfn.{bare}( - bare {bare} becomes #NAME?")
                for quoted, plain in SHEETREF_RE.findall(body):
                    sheet = (quoted or plain).replace("''", "'")
                    if plain and sheet not in names and any(n.endswith(" " + plain) or " " in n and n.split()[-1] == plain for n in names):
                        errors.append(f"{loc}: sheet name with spaces must be quoted: '{[n for n in names if n.endswith(plain)][0]}'!")
                    elif sheet.startswith("[") or sheet in names:
                        pass
                    else:
                        errors.append(f"{loc}: references missing sheet {sheet!r}")
                consts = [m for m in CONST_RE.findall(body) if m not in BENIGN_CONSTANTS]
                if consts:
                    warnings.append(f"{loc}: hard-coded constant(s) {consts} in {v} - move to a labelled input cell")
                m = re.fullmatch(r"=\s*(SUM|AVERAGE|SUBTOTAL\(\d+,)\(?\s*\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)\s*\)", v.replace(" ", ""), re.I)
                if m and m.group(2) == m.group(4) == get_column_letter(c.column):
                    end = int(m.group(5))
                    def counts(val):  # numbers and formulas both feed a total
                        return isinstance(val, (int, float)) and not isinstance(val, bool) or (isinstance(val, str) and val.startswith("="))
                    gap = [r for r in range(end + 1, c.row) if counts(ws.cell(r, c.column).value)]
                    if gap:
                        warnings.append(f"{loc}: {v} stops at row {end} but rows {gap} above the total hold numbers or formulas")
        if no_cache:
            info.append(f"{ws.title}: {no_cache} formula(s) have no cached value (never recalculated) - run recalc.py")
        # a typed-over number sitting inside a column of identical formulas
        for (r, c), f in formulas.items():
            below = formulas.get((r + 2, c))
            mid = ws.cell(r + 1, c).value
            if below and isinstance(mid, (int, float)) and not isinstance(mid, bool) \
                    and signature(f, r, c) == signature(below, r + 2, c):
                warnings.append(f"{ws.title}!{get_column_letter(c)}{r + 1}: hard-coded {mid} between formulas "
                                f"{f} and {below} - restore the formula or label it as an override")
        # pattern breaks along rows and columns
        for axis in ("row", "col"):
            groups = {}
            for (r, c), f in formulas.items():
                groups.setdefault(r if axis == "row" else c, []).append((c if axis == "row" else r, r, c, f))
            for key, cells in groups.items():
                cells.sort()
                run = [cells[0]] if cells else []
                runs = []
                for prev, cur in zip(cells, cells[1:]):
                    if cur[0] == prev[0] + 1:
                        run.append(cur)
                    else:
                        runs.append(run); run = [cur]
                if run:
                    runs.append(run)
                for run in runs:
                    if len(run) < 4:
                        continue
                    sigs = [signature(f, r, c) for _, r, c, f in run]
                    common, n = Counter(sigs).most_common(1)[0]
                    if n >= len(run) - 2 and n < len(run):
                        for (_, r, c, f), s in zip(run, sigs):
                            if s != common:
                                warnings.append(f"{ws.title}!{get_column_letter(c)}{r}: {f} breaks the {axis} pattern of its neighbours")
    return {"file": str(path), "errors": errors, "warnings": sorted(set(warnings)), "info": info}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--sheet")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    rep = audit(Path(a.file), a.sheet)
    if a.json:
        print(json.dumps(rep, indent=2))
    else:
        for k, label in (("errors", "ERROR"), ("warnings", "WARN"), ("info", "INFO")):
            for m in rep[k]:
                print(f"{label}: {m}")
        print(f"{len(rep['errors'])} error(s), {len(rep['warnings'])} warning(s)")
    sys.exit(1 if rep["errors"] else 0)


if __name__ == "__main__":
    main()
