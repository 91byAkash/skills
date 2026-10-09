#!/usr/bin/env python3
"""See a workbook the way you need to edit it: sheets, used ranges, and cells
with formula AND cached value side by side (two openpyxl loads; one load can't
give both).

Usage:
  xlsx_inspect.py book.xlsx                       # overview of every sheet
  xlsx_inspect.py book.xlsx --range "Model!A1:H20"
  xlsx_inspect.py book.xlsx --range "'Q1 Inputs'!B2:B12" --styles   # + font colour/fill/number format
  xlsx_inspect.py book.xlsx --inputs              # cells styled as inputs (blue font or yellow fill)
"""
import argparse, re, sys
from pathlib import Path

import openpyxl
from openpyxl.utils import range_boundaries


def color_of(font_or_fill, attr):
    c = getattr(font_or_fill, attr, None)
    if c is None:
        return None
    rgb = getattr(c, "rgb", None)
    return rgb if isinstance(rgb, str) else None


def overview(wb, wbv):
    for ws in wb.worksheets:
        n_f = sum(1 for r in ws.iter_rows() for c in r if isinstance(c.value, str) and c.value.startswith("="))
        n_v = sum(1 for r in ws.iter_rows() for c in r if c.value is not None)
        print(f"[{ws.title}] {ws.sheet_state} dims={ws.dimensions} cells={n_v} formulas={n_f} "
              f"merged={len(ws.merged_cells.ranges)} freeze={ws.freeze_panes} validations={len(ws.data_validations.dataValidation)} "
              f"cond_formats={len(ws.conditional_formatting)} charts={len(getattr(ws, '_charts', []))}")
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1), [])][:12]
        print("   row 1:", header)
    if wb.defined_names:
        try:
            items = list(wb.defined_names.items())
        except AttributeError:
            items = [(d.name, d) for d in wb.defined_names.definedName]
        for name, dn in items:
            print(f"   name {name} = {dn.attr_text}")


def dump(wb, wbv, spec, styles):
    m = re.match(r"^(?:'((?:[^']|'')+)'|([^!]+))!(.+)$", spec)
    if not m:
        sys.exit("error: --range like Model!A1:H20 or 'Sheet Name'!B2:B9")
    sheet = (m.group(1) or m.group(2)).replace("''", "'")
    ws, vs = wb[sheet], wbv[sheet]
    min_col, min_row, max_col, max_row = range_boundaries(m.group(3))
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for c in row:
            if c.value is None:
                continue
            cached = vs[c.coordinate].value
            line = f"{c.coordinate:>7}  {c.value!r}"
            if isinstance(c.value, str) and c.value.startswith("="):
                line += f"  ->  {cached!r}"
            if styles:
                line += f"   fmt={c.number_format!r} font={color_of(c.font, 'color')} fill={color_of(c.fill, 'fgColor')} bold={c.font.b}"
            print(line)


def inputs(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                fc = (color_of(c.font, "color") or "")[-6:]
                fill = (color_of(c.fill, "fgColor") or "")[-6:] if c.fill and c.fill.fill_type == "solid" else ""
                if fc == "0000FF" or fill == "FFFF00":
                    print(f"{ws.title}!{c.coordinate}: {c.value!r} (font {fc or '-'}, fill {fill or '-'})")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--range")
    ap.add_argument("--styles", action="store_true")
    ap.add_argument("--inputs", action="store_true")
    a = ap.parse_args()
    vba = a.file.lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(a.file, keep_vba=vba)
    wbv = openpyxl.load_workbook(a.file, data_only=True)
    if a.range:
        dump(wb, wbv, a.range, a.styles)
    elif a.inputs:
        inputs(wb)
    else:
        overview(wb, wbv)


if __name__ == "__main__":
    main()
