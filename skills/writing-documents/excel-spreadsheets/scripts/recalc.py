#!/usr/bin/env python3
"""Recalculate every formula with LibreOffice and report errors as JSON.

Usage:
  recalc.py model.xlsx                 # check only: your file is not modified
  recalc.py model.xlsx --write         # also store computed values in the file (for previewers/pandas)
  recalc.py model.xlsx --timeout 90 --max-locations 50
  recalc.py linked.xlsx --force        # proceed although the workbook links to other files

Output (stdout, JSON):
  {"status": "success" | "errors_found", "total_formulas": 412, "total_errors": 2,
   "error_summary": {"#DIV/0!": {"count": 2, "locations": ["Model!F9", ...], "locations_truncated": 0}},
   "suspect_functions": {"Model!C4": "xlookup"}, "written": false}
or {"error": "..."} when nothing could be recalculated.
Exit codes: 0 = recalculated, no errors; 1 = could not recalculate; 2 = recalculated but errors_found.
A clean run proves formulas evaluate, not that they reference the right cells.
"""
import argparse, json, re, shutil, sys, tempfile, zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from office_convert import convert  # noqa: E402

ERRORS = ("#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#SPILL!", "#CALC!", "Err:")
FUNC_RE = re.compile(r"\b([A-Za-z_][A-Za-z0-9_.]*)\s*\(")


def external_links(path):
    with zipfile.ZipFile(path) as z:
        return [n for n in z.namelist() if n.startswith("xl/externalLinks/") and n.endswith(".xml")]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("timeout_pos", nargs="?", type=int, help=argparse.SUPPRESS)
    ap.add_argument("--timeout", type=int, default=60)
    ap.add_argument("--write", action="store_true", help="replace the file with the recalculated copy")
    ap.add_argument("--force", action="store_true", help="recalculate despite external links (their values will be lost)")
    ap.add_argument("--max-locations", type=int, default=100)
    a = ap.parse_args()
    timeout = a.timeout_pos or a.timeout
    src = Path(a.file)
    try:
        import openpyxl
        if not src.exists():
            raise FileNotFoundError(src)
        links = external_links(src)
        if links and not a.force:
            print(json.dumps({"error": f"workbook links to {len(links)} external file(s); LibreOffice cannot resolve them and "
                              "would replace their cached values with #REF!/#NAME?. Copy the needed values in as constants "
                              "first, or pass --force to accept the loss.", "external_links": links}, indent=2))
            sys.exit(1)
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp) / src.name
            shutil.copy(src, work)
            outdir = Path(tmp) / "out"
            out = convert(work, src.suffix.lstrip(".").lower() if src.suffix.lower() in (".xlsx", ".xlsm") else "xlsx",
                          outdir=outdir, timeout=timeout, recalc=True)
            keep_vba = src.suffix.lower() == ".xlsm"
            wb_f = openpyxl.load_workbook(src, keep_vba=keep_vba)
            wb_v = openpyxl.load_workbook(out, data_only=True)
            wb_lo = openpyxl.load_workbook(out)
            total, summary, suspects = 0, {}, {}
            for ws in wb_f.worksheets:
                if ws.title not in wb_v.sheetnames:
                    continue
                vs, ls = wb_v[ws.title], wb_lo[ws.title]
                for row in ws.iter_rows():
                    for c in row:
                        if not (isinstance(c.value, str) and c.value.startswith("=")):
                            continue
                        total += 1
                        v = vs[c.coordinate].value
                        loc = f"{ws.title}!{c.coordinate}"
                        if isinstance(v, str) and v.startswith(ERRORS):
                            key = "Err:" if v.startswith("Err:") else v
                            e = summary.setdefault(key, {"count": 0, "locations": [], "locations_truncated": 0})
                            e["count"] += 1
                            if len(e["locations"]) < a.max_locations:
                                e["locations"].append(loc)
                            else:
                                e["locations_truncated"] += 1
                        lo_f = ls[c.coordinate].value
                        if isinstance(lo_f, str):
                            for fn in FUNC_RE.findall(lo_f):
                                if fn.islower() and fn.upper() in c.value.upper():
                                    suspects[loc] = fn  # LibreOffice lowercases functions it cannot parse
            nerr = sum(e["count"] for e in summary.values())
            report = {"status": "errors_found" if nerr else "success", "total_formulas": total, "total_errors": nerr,
                      "error_summary": summary, "suspect_functions": suspects, "written": False}
            if a.write:
                shutil.copy(out, src)
                report["written"] = True
                report["note"] = "file now holds LibreOffice-computed values; re-open with openpyxl before further edits"
        print(json.dumps(report, indent=2))
    except Exception as e:  # noqa: BLE001 - CLI boundary
        print(json.dumps({"error": str(e)}, indent=2))
        sys.exit(1)
    if report.get("status") == "errors_found":
        sys.exit(2)


if __name__ == "__main__":
    main()
